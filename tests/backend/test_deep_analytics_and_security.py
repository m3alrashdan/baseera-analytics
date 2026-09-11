from __future__ import annotations

import csv
import io
from pathlib import Path

import openpyxl
import pytest
from baseera.analytics import dataset_metric
from baseera.ingestion import parse_upload
from baseera.main import create_app
from baseera.reports import render_report_export
from fastapi.testclient import TestClient
from sqlalchemy import event

from .conftest import login


def test_margin_amount_and_rate_are_distinct_governed_metrics() -> None:
    rows = [
        {"revenue": 100, "cost": 60},
        {"revenue": 50, "cost": 20},
    ]

    amount = dataset_metric("gross_margin", rows)
    rate = dataset_metric("gross_margin_rate", rows)

    assert amount == {"status": "completed", "value": 70.0, "unit": "JOD", "warnings": []}
    assert rate == {
        "status": "completed",
        "value": pytest.approx(46.6666666667),
        "unit": "%",
        "warnings": [],
    }


def test_margin_rate_abstains_for_zero_revenue_or_missing_cost() -> None:
    zero_denominator = dataset_metric(
        "gross_margin_rate",
        [{"revenue": 50, "cost": 20}, {"revenue": -50, "cost": -15}],
    )
    missing_cost = dataset_metric(
        "gross_margin_rate",
        [{"revenue": 50, "cost": None}],
    )

    assert zero_denominator["status"] == "insufficient_data"
    assert zero_denominator["value"] is None
    assert "zero" in zero_denominator["warnings"][0].lower()
    assert missing_cost["status"] == "insufficient_data"
    assert "cost" in missing_cost["warnings"][0].lower()


def test_company_metrics_use_portable_datetime_queries(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    statements: list[str] = []

    def capture_statement(
        _connection: object,
        _cursor: object,
        statement: str,
        _parameters: object,
        _context: object,
        _executemany: bool,
    ) -> None:
        statements.append(statement.lower())

    event.listen(client.app.state.engine, "before_cursor_execute", capture_statement)
    try:
        response = client.get("/api/v1/overview")
    finally:
        event.remove(client.app.state.engine, "before_cursor_execute", capture_statement)

    assert response.status_code == 200
    assert response.json()["metrics"]["avg_resolution_hours"]["value"] > 0
    assert not any("julianday" in statement for statement in statements)


def test_forecast_selects_calibrates_and_evaluates_on_disjoint_windows(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    response = client.post(
        "/api/v1/forecasts",
        json={"metric_id": "order_count", "horizon": 4},
        headers=auth_headers,
    )

    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["status"] == "completed"
    backtest = payload["backtest"]
    assert backtest["strategy"] == "rolling_origin_selection_then_reserved_holdout"
    assert backtest["selection_periods"] >= 3
    assert backtest["evaluation_periods"] >= 2
    assert backtest["wape"] >= 0

    # The MASE denominator is the benchmark's own mean absolute change over the training
    # series. A scale taken from an arbitrary early slice can be near zero and inflates
    # the ratio into the hundreds, which reads as catastrophic but measures nothing.
    assert backtest["mase_scale"] > 0
    assert backtest["mase_scale_basis"] in {"seasonal_naive_12", "naive_one_step"}
    assert 0 <= backtest["mase"] < 10

    # Selection, calibration and evaluation must not share data, or the interval ends up
    # calibrated on exactly the window the winner was chosen for fitting best.
    assert payload["model"]["windows_are_disjoint"] is True
    assert (
        payload["model"]["selected_on_range"]["end"]
        < payload["model"]["evaluated_on_range"]["start"]
    )
    assert payload["model"]["interval_method"] == "conformal_empirical_quantiles_per_horizon"

    assert 0 <= backtest["interval_coverage"] <= 1
    assert backtest["interval_target"] == payload["interval_level"]
    assert len(backtest["candidates"]) >= 5
    assert payload["model"]["name"] in {item["name"] for item in backtest["candidates"]}

    assert len(payload["forecast"]) == 4
    # An interval must contain its own point forecast, and must not narrow with distance.
    widths = []
    for item in payload["forecast"]:
        assert item["lower"] <= item["value"] <= item["upper"]
        widths.append(item["interval_width"])
    assert widths == sorted(widths)

    assert payload["history"], "the history the model was fitted on must be returned"
    assert "seasonality" in payload and "residual_diagnostics" in payload


def test_forecast_refuses_to_fit_across_a_structural_break(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    # Seeded support volume nearly doubles in 2026-04 and holds. Fitting the whole history
    # returns a value that was never true before or after the break, with an interval far
    # too narrow to admit it.
    response = client.post(
        "/api/v1/forecasts",
        json={"metric_id": "support_case_count", "horizon": 6},
        headers=auth_headers,
    )

    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["status"] == "regime_change"
    shift = payload["level_shift"]
    assert shift["changed_at"] == "2026-04"
    assert shift["direction"] == "increase"
    assert shift["relative_change"] > 0.5
    assert shift["periods_since_change"] < 6

    # The published figure must sit at the new level, not between the two regimes.
    assert payload["model"]["name"] == "post_break_level"
    assert all(item["value"] == payload["forecast"][0]["value"] for item in payload["forecast"])
    assert (
        shift["after_median"] * 0.9
        <= payload["forecast"][0]["value"]
        <= (shift["after_median"] * 1.1)
    )

    # No accuracy may be claimed when nothing could be held out.
    assert payload["backtest"]["strategy"] == "not_performed"
    assert payload["backtest"]["mase"] is None
    assert payload["backtest"]["interval_coverage"] is None
    assert payload["decision_required"]["en"] and payload["decision_required"]["ar"]
    assert any("changed level" in warning for warning in payload["warnings"])


def test_forecast_covers_several_metrics_and_refuses_the_rest(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    for metric_id in ("net_revenue", "order_count", "gross_margin"):
        response = client.post(
            "/api/v1/forecasts",
            json={"metric_id": metric_id, "horizon": 3},
            headers=auth_headers,
        )
        assert response.status_code == 201, f"{metric_id}: {response.text}"
        assert response.json()["metric_id"] == metric_id

    refused = client.post(
        "/api/v1/forecasts",
        json={"metric_id": "project_delay_rate", "horizon": 3},
        headers=auth_headers,
    )
    assert refused.status_code == 422
    body = refused.json()["error"]
    assert body["code"] == "forecast_metric_unsupported"
    assert "net_revenue" in body["details"]["supported_metrics"]


def test_optimizer_uses_cp_sat_when_optional_dependency_is_installed(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    pytest.importorskip("ortools.sat.python.cp_model")
    response = client.post(
        "/api/v1/optimizations",
        json={
            "workers": [
                {"id": "w1", "skills": ["support"], "capacity_hours": 8},
                {"id": "w2", "skills": ["support"], "capacity_hours": 8},
            ],
            "tasks": [
                {"id": "t1", "required_skill": "support", "hours": 6},
                {"id": "t2", "required_skill": "support", "hours": 2},
            ],
        },
        headers=auth_headers,
    )

    assert response.status_code == 201
    payload = response.json()
    assert payload["feasibility"] == "feasible"
    assert payload["solver"]["name"] == "ortools_cp_sat"
    assert payload["solver"]["status"] == "optimal"
    assert payload["solver"]["optimality_proven"] is True
    assert sum(item["hours"] for item in payload["assignments"]) == 8


def test_causal_randomized_example_is_bounded_reproducible_and_assumption_labeled(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    observations = [
        {"policy_change": 0, "resolution_hours": 10 + index % 3} for index in range(12)
    ] + [{"policy_change": 1, "resolution_hours": 15 + index % 3} for index in range(12)]
    request = {
        "treatment": "policy_change",
        "outcome": "resolution_hours",
        "time_ordering": "treatment_before_outcome",
        "assumptions": ["random_assignment", "consistency", "no_interference"],
        "identification_strategy": "randomized_difference_in_means",
        "observations": observations,
    }

    first = client.post("/api/v1/causal/analyze", json=request, headers=auth_headers)
    second = client.post("/api/v1/causal/analyze", json=request, headers=auth_headers)

    assert first.status_code == 200, first.text
    assert first.json() == second.json()
    payload = first.json()
    assert payload["status"] == "completed"
    assert payload["classification"] == "causal_estimate_under_assumptions"
    assert payload["estimate"]["average_treatment_effect"] == pytest.approx(5.0)
    assert payload["estimate"]["lower_95"] < 5 < payload["estimate"]["upper_95"]
    assert payload["diagnostics"]["observations"] == 24
    assert payload["assumption_status"] == "declared_not_verified"
    assert payload["limitations"]


def test_spreadsheet_exports_neutralize_formula_and_url_interpretation() -> None:
    report = {
        "title": '=HYPERLINK("https://attacker.invalid","open")',
        "version": 1,
        "language": "en",
        "sections": [
            {
                "title": "@SUM(1+1)",
                "kind": "+cmd",
                "content": "ignored",
            }
        ],
    }

    csv_rows = list(
        csv.DictReader(io.StringIO(render_report_export(report, "csv").decode("utf-8-sig")))
    )
    assert csv_rows[0]["section"].startswith("'@")
    assert csv_rows[0]["kind"].startswith("'+")

    workbook = openpyxl.load_workbook(
        io.BytesIO(render_report_export(report, "xlsx")), data_only=False
    )
    sheet = workbook["Report data"]
    assert sheet["A1"].data_type == "s"
    assert sheet["A1"].value.startswith("'=")
    assert sheet["A4"].data_type == "s"
    assert sheet["A4"].value.startswith("'@")
    workbook.close()


def test_xlsx_ingestion_reports_sheet_selection_and_unevaluated_formulas() -> None:
    workbook = openpyxl.Workbook()
    orders = workbook.active
    orders.title = "Orders"
    orders.append(["id", "amount"])
    orders.append(["0001", "=1+1"])
    hidden = workbook.create_sheet("Hidden calculation")
    hidden.sheet_state = "hidden"
    hidden.append(["id", "amount"])
    hidden.append(["0002", 99])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()

    _, parsed = parse_upload("formula.xlsx", buffer.getvalue())

    assert parsed.extraction["selected_sheet"] == "Orders"
    assert parsed.extraction["sheet_selection_policy"] == "first_visible_non_empty"
    assert parsed.extraction["formulas_evaluated"] is False
    assert parsed.extraction["formula_cells_detected"] == 1
    assert parsed.extraction["formula_cells_without_cached_value"] == 1
    assert parsed.extraction["sheet_inventory"][1]["state"] == "hidden"
    assert parsed.rows == [{"id": "0001", "amount": None}]


def test_viewer_cannot_mutate_artifacts_or_read_people_details(client: TestClient) -> None:
    viewer_headers = login(client, "viewer@empty.baseera.local", "BaseeraEmpty!2026")

    people = client.get("/api/v1/company/people")
    upload = client.post(
        "/api/v1/datasets/upload",
        files={"file": ("orders.csv", b"id,revenue\n1,2\n", "text/csv")},
        headers=viewer_headers,
    )
    decision = client.post(
        "/api/v1/decisions", json={"title": "Must not be created"}, headers=viewer_headers
    )

    assert people.status_code == 403
    assert people.json()["error"]["code"] == "permission_denied"
    assert upload.status_code == 403
    assert decision.status_code == 403


def test_people_projection_and_department_scoping_are_enforced(client: TestClient) -> None:
    login(client)
    executive_people = client.get("/api/v1/company/people")
    assert executive_people.status_code == 200
    assert executive_people.json()["items"]
    assert all("headcount" in item for item in executive_people.json()["items"])
    assert all(
        "name" not in item and "skills" not in item for item in executive_people.json()["items"]
    )

    client.post(
        "/api/v1/auth/logout",
        headers={"X-CSRF-Token": client.get("/api/v1/auth/context").json().get("csrf_token", "")},
    )
    manager_headers = login(client, "manager@demo.baseera.local", "BaseeraManager!2026")
    operations = client.get("/api/v1/company/operations")
    costs = client.get("/api/v1/company/costs")
    assistant = client.post(
        "/api/v1/assistant/query",
        json={"question": "What is scoped revenue?", "metric_id": "net_revenue"},
        headers=manager_headers,
    )

    assert operations.status_code == 200
    assert operations.json()["items"] == []
    assert {item["department_id"] for item in costs.json()["items"]} == {"dept-support"}
    assert assistant.status_code == 200
    assert assistant.json()["status"] == "insufficient_data"


def test_origin_policy_and_fixed_window_rate_limits(tmp_path: Path) -> None:
    app = create_app(
        database_url=f"sqlite:///{tmp_path / 'limits.db'}",
        artifact_root=tmp_path / "artifacts",
        testing=True,
        seed_demo=True,
        login_rate_limit=2,
        api_rate_limit=2,
        rate_limit_window_seconds=60,
    )
    with TestClient(app) as client:
        cross_origin = client.post(
            "/api/v1/auth/login",
            json={"email": "nobody@example.test", "password": "wrong"},
            headers={"Origin": "https://attacker.invalid"},
        )
        assert cross_origin.status_code == 403
        assert cross_origin.json()["error"]["code"] == "origin_denied"

        first = client.post(
            "/api/v1/auth/login",
            json={"email": "nobody@example.test", "password": "wrong"},
        )
        second = client.post(
            "/api/v1/auth/login",
            json={"email": "nobody@example.test", "password": "wrong"},
        )
        limited = client.post(
            "/api/v1/auth/login",
            json={"email": "nobody@example.test", "password": "wrong"},
        )
        assert first.status_code == second.status_code == 401
        assert limited.status_code == 429
        assert limited.headers["retry-after"]
        assert limited.json()["error"]["code"] == "rate_limit_exceeded"

        assert client.get("/api/v1/health/live").status_code == 200
        assert client.get("/api/v1/health/live").status_code == 200
        api_limited = client.get("/api/v1/health/live")
        assert api_limited.status_code == 429
        assert api_limited.headers["x-ratelimit-limit"] == "2"
