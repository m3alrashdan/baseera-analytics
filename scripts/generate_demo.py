#!/usr/bin/env python3
"""Generate deterministic, fictional BASEERA ingestion fixtures at a bounded scale.

The output is deliberately synthetic and is never loaded into a database by this
script.  It exercises file parsing, dirty-data review, and scale measurements; it
does not demonstrate real-company accuracy or production capacity.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import sys
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any

REPORTING_CUTOFF = date(2026, 6, 30)
DEPARTMENTS = (
    ("dept-sales", "Sales", "المبيعات"),
    ("dept-support", "Customer Support", "دعم العملاء"),
    ("dept-operations", "Operations", "العمليات"),
    ("dept-technology", "Technology", "التقنية"),
    ("dept-finance", "Finance", "المالية"),
    ("dept-people", "People", "الموارد البشرية"),
    ("dept-strategy", "Strategy", "الاستراتيجية"),
)
OUTPUT_FILES = (
    "departments.csv",
    "employees.csv",
    "customers.csv",
    "projects.csv",
    "sales_orders.csv",
    "expenses.tsv",
    "support_events.jsonl",
    "dirty_sales.xlsx",
    "manifest.json",
)


@dataclass(frozen=True)
class Sizes:
    employees: int
    customers: int
    projects: int
    orders: int
    expenses: int
    support_cases: int


def sizes_for(scale: int) -> Sizes:
    return Sizes(
        employees=100 * scale,
        customers=240 * scale,
        projects=20 * scale,
        orders=960 * scale,
        expenses=360 * scale,
        support_cases=420 * scale,
    )


def iso_day(index: int, span_days: int = 730) -> str:
    return (REPORTING_CUTOFF - timedelta(days=index % span_days)).isoformat()


def write_delimited(
    path: Path, fieldnames: list[str], rows: Iterable[dict[str, Any]], delimiter: str = ","
) -> int:
    count = 0
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter=delimiter)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)
            count += 1
    return count


def employee_rows(rng: random.Random, size: int) -> Iterable[dict[str, Any]]:
    skills = ("analytics|arabic", "support|arabic", "delivery|technical", "finance|planning")
    for index in range(1, size + 1):
        department = DEPARTMENTS[(index - 1) % len(DEPARTMENTS)][0]
        workload = rng.uniform(24, 39)
        if department == "dept-support" and index % 5 == 0:
            workload = rng.uniform(40, 48)
        yield {
            "employee_id": f"{index:06d}",
            "department_id": department,
            "display_name": f"Fictional Colleague {index:06d}",
            "role": "Team lead" if index % 13 == 0 else "Specialist",
            "skills": skills[(index - 1) % len(skills)],
            "capacity_hours": "40.0",
            "workload_hours": f"{workload:.1f}",
        }


def customer_rows(size: int) -> Iterable[dict[str, Any]]:
    segments = ("Enterprise", "Mid-market", "Small business")
    regions = ("Jordan", "Saudi Arabia", "UAE", "Iraq")
    for index in range(1, size + 1):
        yield {
            "customer_id": f"C{index:07d}",
            "name": f"Fictional Customer {index:07d}",
            "segment": segments[(index - 1) % len(segments)],
            "region": regions[(index - 1) % len(regions)],
        }


def project_rows(rng: random.Random, size: int) -> Iterable[dict[str, Any]]:
    for index in range(1, size + 1):
        delayed = index % 5 == 0
        progress = round(rng.uniform(0.25, 0.92), 2)
        budget = 35_000 + index * 1_250
        actual = round(budget * (rng.uniform(0.65, 1.14) if delayed else rng.uniform(0.4, 0.9)), 2)
        yield {
            "project_id": f"P{index:05d}",
            "department_id": DEPARTMENTS[(index - 1) % len(DEPARTMENTS)][0],
            "supplier_id": f"S{index % 6 + 1:03d}" if index % 3 == 0 else "",
            "name": f"Fictional Transformation Initiative {index:05d}",
            "status": "delayed" if delayed else "active",
            "progress": f"{progress:.2f}",
            "budget_jod": f"{budget:.2f}",
            "actual_cost_jod": f"{actual:.2f}",
            "due_date": iso_day(-index, 540),
            "delay_days": str(7 + index % 31 if delayed else 0),
        }


def sales_rows(rng: random.Random, size: int, customers: int) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    regions = ("Jordan", "Saudi Arabia", "UAE", "Iraq")
    for index in range(1, size + 1):
        gross = round(500 + (index % 17) * 83 + rng.uniform(0, 140), 2)
        discount = round(gross * (0.17 if index % 4 == 0 else 0.04), 2)
        status = "return" if index % 113 == 0 else "invoiced"
        sign = -1 if status == "return" else 1
        net = round(sign * (gross - discount), 2)
        cost: str | float = round(sign * abs(net) * rng.uniform(0.58, 0.73), 2)
        if index % 97 == 0:
            cost = ""
        rows.append(
            {
                "order_id": f"{index:08d}",
                "batch_id": f"B{(index - 1) // 80 + 1:05d}",
                "order_date": iso_day(size - index),
                "customer_id": f"C{(index - 1) % customers + 1:07d}",
                "region": regions[(index - 1) % len(regions)],
                "status": status,
                "gross_jod": gross,
                "discount_jod": discount,
                "net_revenue_jod": net,
                "cost_jod": cost,
            }
        )
    return rows


def expense_rows(rng: random.Random, size: int) -> Iterable[dict[str, Any]]:
    categories = ("hosting", "travel", "equipment", "training", "facilities")
    for index in range(1, size + 1):
        currency = "USD" if index % 19 == 0 else "JOD"
        exchange_rate: str | float = 0.709 if currency == "USD" else 1
        if index % 171 == 0:
            exchange_rate = ""
        yield {
            "expense_id": f"E{index:07d}",
            "department_id": DEPARTMENTS[(index - 1) % len(DEPARTMENTS)][0],
            "expense_date": iso_day(size - index),
            "category": categories[(index - 1) % len(categories)],
            "amount": f"{rng.uniform(90, 5400):.2f}",
            "currency": currency,
            "jod_exchange_rate": exchange_rate,
        }


def write_events(path: Path, rng: random.Random, cases: int) -> int:
    event_count = 0
    with path.open("w", encoding="utf-8") as handle:
        for case_index in range(1, cases + 1):
            opened = datetime(2026, 1, 1, 8, tzinfo=UTC) + timedelta(hours=case_index * 5)
            wait_hours = rng.randint(1, 20)
            work_hours = rng.randint(1, 12)
            events = [
                ("opened", "start", opened),
                ("assigned", "complete", opened + timedelta(hours=wait_hours)),
                ("resolved", "complete", opened + timedelta(hours=wait_hours + work_hours)),
            ]
            if case_index % 17 == 0:
                events.append(
                    (
                        "reopened",
                        "complete",
                        opened + timedelta(hours=wait_hours + work_hours + 6),
                    )
                )
            for activity, lifecycle, occurred_at in events:
                event_count += 1
                record = {
                    "event_id": f"EV{event_count:09d}",
                    "case_id": f"CASE{case_index:07d}",
                    "activity": activity,
                    "lifecycle": lifecycle,
                    "occurred_at": occurred_at.isoformat().replace("+00:00", "Z"),
                    "resource_ref": f"resource-{case_index % 23:03d}",
                    "source": "synthetic-generator",
                }
                handle.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")
    return event_count


def write_dirty_workbook(path: Path, sales: list[dict[str, Any]]) -> int:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required; run `uv sync --frozen` first") from exc

    workbook = Workbook()
    workbook.properties.creator = "BASEERA synthetic fixture generator"
    workbook.properties.created = datetime(2026, 6, 30, 0, 0)
    workbook.properties.modified = datetime(2026, 6, 30, 0, 0)
    sheet = workbook.active
    sheet.title = "مبيعات Sales"
    headers = list(sales[0])
    sheet.append(headers)
    selected = [dict(row) for row in sales[: min(500, len(sales))]]
    if selected:
        selected.append(dict(selected[min(11, len(selected) - 1)]))
        selected[-1]["batch_id"] = selected[min(11, len(selected) - 1)]["batch_id"]
    if len(selected) > 4:
        selected[1]["order_date"] = "01/02/2026"
        selected[2]["net_revenue_jod"] = "1.234,50"
        selected[3]["order_date"] = "31/02/2026"
        selected[4]["cost_jod"] = ""
    for row in selected:
        sheet.append([row[field] for field in headers])
    order_column = headers.index("order_id") + 1
    for cell in sheet.iter_cols(min_col=order_column, max_col=order_column, min_row=2):
        for item in cell:
            item.number_format = "@"

    notes = workbook.create_sheet("README")
    notes.append(["Synthetic fixture", "ملف تجريبي اصطناعي"])
    notes.append(["Reporting cutoff", REPORTING_CUTOFF.isoformat()])
    notes.append(["Known conditions", "duplicate row; leading-zero IDs; valid negative returns"])
    notes.append(["Review required", "missing cost; ambiguous locale; invalid date; mixed decimal"])
    notes.append(["Formula limitation", "=SUM('مبيعات Sales'!I2:I20)"])
    workbook.save(path)
    canonical = path.with_suffix(".canonical.xlsx")
    with (
        zipfile.ZipFile(path, "r") as source,
        zipfile.ZipFile(
            canonical, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as target,
    ):
        for name in sorted(source.namelist()):
            source_info = source.getinfo(name)
            info = zipfile.ZipInfo(name, date_time=(2026, 6, 30, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = source_info.external_attr
            info.create_system = source_info.create_system
            target.writestr(info, source.read(name))
    canonical.replace(path)
    return len(selected)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def prepare_output(path: Path, force: bool) -> Path:
    resolved = path.expanduser().resolve()
    if resolved == Path(resolved.anchor):
        raise ValueError("refusing to use a filesystem root as output")
    resolved.mkdir(parents=True, exist_ok=True)
    conflicts = [resolved / name for name in OUTPUT_FILES if (resolved / name).exists()]
    if conflicts and not force:
        rendered = ", ".join(item.name for item in conflicts)
        raise ValueError(f"output already contains generated files ({rendered}); use --force")
    return resolved


def generate(output: Path, scale: int, seed: int, force: bool) -> dict[str, Any]:
    output = prepare_output(output, force)
    rng = random.Random(seed)
    sizes = sizes_for(scale)
    counts: dict[str, int] = {}

    counts["departments.csv"] = write_delimited(
        output / "departments.csv",
        ["department_id", "name_en", "name_ar"],
        (
            {"department_id": identifier, "name_en": name_en, "name_ar": name_ar}
            for identifier, name_en, name_ar in DEPARTMENTS
        ),
    )
    counts["employees.csv"] = write_delimited(
        output / "employees.csv",
        [
            "employee_id",
            "department_id",
            "display_name",
            "role",
            "skills",
            "capacity_hours",
            "workload_hours",
        ],
        employee_rows(rng, sizes.employees),
    )
    counts["customers.csv"] = write_delimited(
        output / "customers.csv",
        ["customer_id", "name", "segment", "region"],
        customer_rows(sizes.customers),
    )
    counts["projects.csv"] = write_delimited(
        output / "projects.csv",
        [
            "project_id",
            "department_id",
            "supplier_id",
            "name",
            "status",
            "progress",
            "budget_jod",
            "actual_cost_jod",
            "due_date",
            "delay_days",
        ],
        project_rows(rng, sizes.projects),
    )
    sales = sales_rows(rng, sizes.orders, sizes.customers)
    counts["sales_orders.csv"] = write_delimited(output / "sales_orders.csv", list(sales[0]), sales)
    counts["expenses.tsv"] = write_delimited(
        output / "expenses.tsv",
        [
            "expense_id",
            "department_id",
            "expense_date",
            "category",
            "amount",
            "currency",
            "jod_exchange_rate",
        ],
        expense_rows(rng, sizes.expenses),
        delimiter="\t",
    )
    counts["support_events.jsonl"] = write_events(
        output / "support_events.jsonl", rng, sizes.support_cases
    )
    counts["dirty_sales.xlsx"] = write_dirty_workbook(output / "dirty_sales.xlsx", sales)

    artifacts = []
    for name in OUTPUT_FILES:
        if name == "manifest.json":
            continue
        path = output / name
        artifacts.append(
            {
                "path": name,
                "records": counts[name],
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
        )
    manifest = {
        "schemaVersion": "1.0",
        "generator": "scripts/generate_demo.py",
        "generatedAt": datetime.now(UTC).isoformat().replace("+00:00", "Z"),
        "seed": seed,
        "scale": scale,
        "fixture": {
            "fictional": True,
            "tenantId": "tenant-demo-generated",
            "reportingCutoff": REPORTING_CUTOFF.isoformat(),
            "timezone": "Asia/Amman",
            "currency": "JOD",
        },
        "knownDirtyConditions": [
            "duplicate_import_row",
            "leading_zero_identifier",
            "legitimate_negative_return",
            "missing_cost",
            "ambiguous_day_month_date",
            "invalid_date",
            "mixed_decimal_locale",
            "spreadsheet_formula_without_engine_recalculation",
        ],
        "artifacts": artifacts,
        "limitations": [
            "Synthetic data validates implementation behavior only.",
            (
                "The generation timestamp is intentionally non-deterministic; data rows and "
                "hashes are deterministic for a fixed tool version, seed, and scale."
            ),
            (
                "Performance conclusions require a separately recorded machine, resource, and "
                "concurrency profile."
            ),
        ],
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("var/generated-demo"))
    parser.add_argument("--scale", type=int, default=1, help="integer scale from 1 to 25")
    parser.add_argument("--seed", type=int, default=20260630)
    parser.add_argument("--force", action="store_true", help="overwrite only known generated files")
    args = parser.parse_args()
    if not 1 <= args.scale <= 25:
        parser.error("--scale must be between 1 and 25")
    try:
        manifest = generate(args.output, args.scale, args.seed, args.force)
    except (OSError, RuntimeError, ValueError) as exc:
        print(f"generation failed: {exc}", file=sys.stderr)
        return 1
    total_records = sum(item["records"] for item in manifest["artifacts"])
    print(
        json.dumps(
            {
                "status": "generated",
                "output": str(args.output.resolve()),
                "artifacts": len(manifest["artifacts"]),
                "records": total_records,
                "seed": args.seed,
                "scale": args.scale,
                "warning": "Synthetic data is not evidence of real-world accuracy or scale.",
            },
            indent=2,
            sort_keys=True,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
