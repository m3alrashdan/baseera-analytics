"""The report writer's deliverables: the analysis dossier as HTML, PDF, Word,
PowerPoint (with native, editable charts), Markdown and JSON, in Arabic or English.

Charts in HTML/PDF are rendered server-side as inline SVG so the documents are
self-contained: no scripts, no network fetches, nothing to break when forwarded.
"""

from __future__ import annotations

import html
import io
import json
import math
from datetime import UTC, datetime
from typing import Any

from ..errors import AppError
from .common import fmt

MEDIA_TYPES = {
    "html": ("text/html; charset=utf-8", "html"),
    "pdf": ("application/pdf", "pdf"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "pptx": ("application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx"),
    "md": ("text/markdown; charset=utf-8", "md"),
    "json": ("application/json", "json"),
}
PALETTE = ["#087f8c", "#c2410c", "#4f46e5", "#15803d", "#b45309", "#be185d", "#0369a1", "#6d28d9"]
CONFIDENCE = {
    "high": {"en": "High confidence", "ar": "ثقة عالية"},
    "medium": {"en": "Medium confidence", "ar": "ثقة متوسطة"},
    "low": {"en": "Low confidence", "ar": "ثقة منخفضة"},
}
T = {
    "title": {"en": "Analysis dossier", "ar": "تقرير التحليل"},
    "prepared_by": {
        "en": "Prepared by the BASEERA AI analyst team",
        "ar": "أعدّه فريق محللي بصيرة الأذكياء",
    },
    "summary": {"en": "Executive summary", "ar": "الملخص التنفيذي"},
    "recommendations": {"en": "Recommendations", "ar": "التوصيات"},
    "findings": {"en": "Findings", "ar": "النتائج"},
    "next": {"en": "Questions worth asking next", "ar": "أسئلة جديرة بالطرح لاحقًا"},
    "method": {"en": "Method, verification and limits", "ar": "المنهجية والتحقق والحدود"},
    "actions": {"en": "Actions", "ar": "الإجراءات"},
    "impact": {"en": "Expected impact", "ar": "الأثر المتوقع"},
    "priority": {"en": "Priority", "ar": "الأولوية"},
    "effort": {"en": "Effort", "ar": "الجهد"},
    "caveats": {"en": "Caveats", "ar": "تحفظات"},
    "so_what": {"en": "So what", "ar": "ماذا يعني ذلك"},
    "engine": {"en": "Engine", "ar": "المحرك"},
    "verification": {
        "en": "Figures verified against evidence",
        "ar": "الأرقام متحقق منها مقابل الأدلة",
    },
    "unverified": {"en": "Figures that could not be traced", "ar": "أرقام تعذّر تتبعها"},
    "deterministic": {
        "en": "All figures were computed by deterministic, reproducible tools.",
        "ar": "جميع الأرقام محسوبة بأدوات حتمية قابلة لإعادة الإنتاج.",
    },
    "risks": {"en": "Risks and caveats", "ar": "المخاطر والتحفظات"},
    "effort_low": {"en": "Low", "ar": "منخفض"},
    "effort_medium": {"en": "Medium", "ar": "متوسط"},
    "effort_high": {"en": "High", "ar": "مرتفع"},
    "evidence": {"en": "Evidence", "ar": "الدليل"},
}


def _t(key: str, locale: str) -> str:
    return T[key][locale]


def _loc(value: Any, locale: str) -> str:
    if isinstance(value, dict):
        return str(value.get(locale) or value.get("en") or value.get("ar") or "")
    return "" if value is None else str(value)


def _metric_value(item: dict[str, Any]) -> str:
    value = item.get("value")
    kind = item.get("format")
    if value is None:
        return "—"
    if kind == "percent" and isinstance(value, int | float):
        return f"{value * 100:.1f}%"
    if kind == "ratio" and isinstance(value, int | float):
        return f"{value:.2f}"
    if kind == "score" and isinstance(value, int | float):
        return f"{value:.0f}/100"
    if isinstance(value, int | float):
        return fmt(value)
    return str(value)


# --------------------------------------------------------------------------- SVG charts
def svg_chart(
    chart: dict[str, Any] | None, locale: str, width: int = 640, height: int = 260
) -> str:
    if not chart:
        return ""
    kind = chart.get("type")
    try:
        if kind in {"line", "forecast"}:
            return _svg_line(chart, locale, width, height)
        if kind in {"bar", "pareto", "bar_horizontal"}:
            return _svg_bar(chart, locale, width, height)
        if kind == "waterfall":
            return _svg_waterfall(chart, locale, width, height)
    except (TypeError, ValueError, ZeroDivisionError):
        return ""
    return ""


def _scale(values: list[float], low: float, high: float) -> tuple[float, float]:
    finite = [v for v in values if isinstance(v, int | float) and not math.isnan(v)]
    if not finite:
        return 0.0, 1.0
    vmin, vmax = min(finite + [0.0]), max(finite)
    if vmax == vmin:
        vmax = vmin + 1
    return vmin, vmax


def _svg_frame(title: str, body: str, width: int, height: int, locale: str) -> str:
    direction = "rtl" if locale == "ar" else "ltr"
    return (
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height + 28}" '
        f'width="{width}" height="{height + 28}" style="display:block;max-width:100%;height:auto" '
        f'role="img" aria-label="{html.escape(title)}" direction="{direction}">'
        f'<text x="{width / 2}" y="16" text-anchor="middle" font-size="13" fill="#142536" '
        f'font-weight="600">{html.escape(title)}</text>'
        f'<g transform="translate(0,28)">{body}</g></svg>'
    )


def _axis_labels(labels: list[str], x_of: Any, y: float, every: int) -> str:
    return "".join(
        f'<text x="{x_of(i):.1f}" y="{y:.1f}" font-size="9" fill="#52697e" text-anchor="middle">'
        f"{html.escape(str(label)[:12])}</text>"
        for i, label in enumerate(labels)
        if i % every == 0
    )


def _svg_line(chart: dict[str, Any], locale: str, width: int, height: int) -> str:
    left, right, top, bottom = 56, 12, 8, 26
    x_labels = [str(x) for x in chart.get("x", [])]
    if chart.get("type") == "forecast":
        history = chart.get("history", [])
        future = chart.get("forecast", [])
        series = [
            {"name": "history", "data": history + [None] * len(future)},
            {
                "name": "forecast",
                "data": [None] * (len(history) - 1) + [history[-1] if history else None] + future,
            },
        ]
        band_low = [None] * len(history) + chart.get("lower", [])
        band_high = [None] * len(history) + chart.get("upper", [])
    else:
        series = chart.get("series", [])[:6]
        band_low = band_high = []
    values = [v for s in series for v in s.get("data", []) if isinstance(v, int | float)]
    values += [v for v in band_low + band_high if isinstance(v, int | float)]
    vmin, vmax = _scale(values, 0, 1)
    n = max(1, len(x_labels) - 1)
    plot_w, plot_h = width - left - right, height - top - bottom

    def x_of(i: int) -> float:
        return left + plot_w * i / n

    def y_of(v: float) -> float:
        return top + plot_h * (1 - (v - vmin) / (vmax - vmin))

    parts = [
        f'<line x1="{left}" y1="{top + plot_h}" x2="{width - right}" y2="{top + plot_h}" stroke="#c8d0d0"/>'
    ]
    for tick in range(5):
        value = vmin + (vmax - vmin) * tick / 4
        y = y_of(value)
        parts.append(
            f'<line x1="{left}" y1="{y:.1f}" x2="{width - right}" y2="{y:.1f}" stroke="#eef1f0"/>'
            f'<text x="{left - 6}" y="{y + 3:.1f}" font-size="9" fill="#52697e" text-anchor="end">{fmt(value)}</text>'
        )
    if band_low:
        points_high = [
            (x_of(i), y_of(v)) for i, v in enumerate(band_high) if isinstance(v, int | float)
        ]
        points_low = [
            (x_of(i), y_of(v)) for i, v in enumerate(band_low) if isinstance(v, int | float)
        ]
        if points_high:
            polygon = " ".join(f"{x:.1f},{y:.1f}" for x, y in points_high + points_low[::-1])
            parts.append(f'<polygon points="{polygon}" fill="#087f8c" fill-opacity="0.15"/>')
    for index, item in enumerate(series):
        color = PALETTE[index % len(PALETTE)]
        points = [
            f"{x_of(i):.1f},{y_of(v):.1f}"
            for i, v in enumerate(item.get("data", []))
            if isinstance(v, int | float)
        ]
        dash = (
            ' stroke-dasharray="5 4"'
            if item.get("style") == "dashed" or item.get("name") == "forecast"
            else ""
        )
        if points:
            parts.append(
                f'<polyline points="{" ".join(points)}" fill="none" stroke="{color}" stroke-width="2"{dash}/>'
            )
    every = max(1, len(x_labels) // 10)
    parts.append(_axis_labels(x_labels, x_of, top + plot_h + 16, every))
    return _svg_frame(_loc(chart.get("title"), locale), "".join(parts), width, height, locale)


def _svg_bar(chart: dict[str, Any], locale: str, width: int, height: int) -> str:
    raw_labels = chart.get("x_ar") if locale == "ar" and chart.get("x_ar") else chart.get("x")
    labels = [str(label) for label in raw_labels or []][:20]
    data = [
        v if isinstance(v, int | float) else 0.0
        for v in (chart.get("series") or [{}])[0].get("data", [])
    ][: len(labels)]
    if not data:
        return ""
    horizontal = chart.get("type") == "bar_horizontal"
    vmin, vmax = _scale(data, 0, 1)
    parts: list[str] = []
    if horizontal:
        left, top = 150, 6
        row = (height - top) / max(1, len(data))
        span = max(abs(vmax), abs(vmin)) or 1
        for i, (label, value) in enumerate(zip(labels, data, strict=False)):
            w = (width - left - 60) * abs(value) / span
            y = top + i * row
            parts.append(
                f'<text x="{left - 6}" y="{y + row * 0.62:.1f}" font-size="10" fill="#142536" text-anchor="end">{html.escape(label[:24])}</text>'
                f'<rect x="{left}" y="{y + row * 0.15:.1f}" width="{w:.1f}" height="{row * 0.7:.1f}" rx="3" fill="{PALETTE[0]}"/>'
                f'<text x="{left + w + 4:.1f}" y="{y + row * 0.62:.1f}" font-size="9" fill="#52697e">{fmt(value * 100) + "%" if abs(value) <= 1 else fmt(value)}</text>'
            )
        return _svg_frame(_loc(chart.get("title"), locale), "".join(parts), width, height, locale)
    left, right, top, bottom = 56, 12, 8, 34
    plot_w, plot_h = width - left - right, height - top - bottom
    step = plot_w / len(data)

    def y_of(v: float) -> float:
        return top + plot_h * (1 - (v - vmin) / (vmax - vmin))

    zero = y_of(0.0)
    for i, value in enumerate(data):
        y = y_of(value)
        parts.append(
            f'<rect x="{left + i * step + step * 0.15:.1f}" y="{min(y, zero):.1f}" width="{step * 0.7:.1f}" '
            f'height="{abs(zero - y):.1f}" rx="2" fill="{PALETTE[0] if value >= 0 else PALETTE[1]}"/>'
        )
    for tick in range(5):
        value = vmin + (vmax - vmin) * tick / 4
        parts.append(
            f'<text x="{left - 6}" y="{y_of(value) + 3:.1f}" font-size="9" fill="#52697e" text-anchor="end">{fmt(value)}</text>'
        )
    every = max(1, len(labels) // 12)
    parts.append(
        _axis_labels(labels, lambda i: left + i * step + step / 2, top + plot_h + 14, every)
    )
    reference = chart.get("reference")
    if isinstance(reference, int | float) and vmin <= reference <= vmax:
        parts.append(
            f'<line x1="{left}" x2="{width - right}" y1="{y_of(reference):.1f}" y2="{y_of(reference):.1f}" stroke="#99620a" stroke-dasharray="4 3"/>'
        )
    return _svg_frame(_loc(chart.get("title"), locale), "".join(parts), width, height, locale)


def _svg_waterfall(chart: dict[str, Any], locale: str, width: int, height: int) -> str:
    labels = [str(x) for x in chart.get("x", [])]
    start, end = float(chart.get("start") or 0), float(chart.get("end") or 0)
    deltas = [float(d) for d in chart.get("deltas", [])]
    levels = [start]
    running = start
    for delta in deltas:
        running += delta
        levels.append(running)
    values = levels + [end]
    vmin, vmax = _scale(values, 0, 1)
    left, right, top, bottom = 56, 12, 8, 34
    plot_w, plot_h = width - left - right, height - top - bottom
    count = len(deltas) + 2
    step = plot_w / count

    def y_of(v: float) -> float:
        return top + plot_h * (1 - (v - vmin) / (vmax - vmin))

    parts = []
    bars = [(start, 0.0, "#385067")]
    running = start
    for delta in deltas:
        bars.append((running + delta, running, PALETTE[3] if delta >= 0 else PALETTE[1]))
        running += delta
    bars.append((end, 0.0, "#385067"))
    for i, (a, b, color) in enumerate(bars):
        y1, y2 = y_of(max(a, b)), y_of(min(a, b))
        parts.append(
            f'<rect x="{left + i * step + step * 0.15:.1f}" y="{y1:.1f}" width="{step * 0.7:.1f}" height="{max(1.0, y2 - y1):.1f}" fill="{color}" rx="2"/>'
        )
    parts.append(
        _axis_labels(labels[:count], lambda i: left + i * step + step / 2, top + plot_h + 14, 1)
    )
    return _svg_frame(_loc(chart.get("title"), locale), "".join(parts), width, height, locale)


# ------------------------------------------------------------------------------- HTML
def dossier_html(dossier: dict[str, Any], locale: str, for_pdf: bool = False) -> str:
    esc = html.escape
    rtl = locale == "ar"
    dataset = dossier.get("dataset", {})
    findings = {f["id"]: f for f in dossier.get("findings", [])}
    summary = _loc(dossier.get("executive_summary"), locale)
    headline = _loc(dossier.get("headline"), locale) if dossier.get("headline") else ""
    engine = dossier.get("engine", {})
    verification = dossier.get("verification", {})

    def paragraphs(text: str) -> str:
        return "".join(f"<p>{esc(block)}</p>" for block in text.split("\n\n") if block.strip())

    recs = []
    for rec in dossier.get("recommendations", []):
        actions = (rec.get("actions_model") or {}).get(locale) or [
            _loc(a, locale) for a in rec.get("actions", [])
        ]
        recs.append(
            f'<div class="rec"><div class="rec-head"><span class="badge p{rec["priority"][-1]}">{esc(rec["priority"])}</span>'
            f"<h3>{esc(_loc(rec['title'], locale))}</h3></div>"
            f"<p>{esc(_loc(rec['rationale'], locale))}</p>"
            f'<p class="impact"><strong>{_t("impact", locale)}:</strong> {esc(_loc(rec["expected_impact"], locale))}</p>'
            f"<ul>{''.join(f'<li>{esc(a)}</li>' for a in actions)}</ul>"
            f'<p class="meta">{_t("effort", locale)}: {_t("effort_" + rec.get("effort", "medium"), locale)}</p></div>'
        )
    sections = []
    for section in dossier.get("sections", []):
        items = []
        for finding_id in section["findings"]:
            f = findings.get(finding_id)
            if not f:
                continue
            metrics = "".join(
                f'<div class="metric"><span>{esc(_loc(m.get("label"), locale))}</span><strong>{esc(_metric_value(m))}</strong></div>'
                for m in f.get("metrics", [])[:4]
            )
            so_what = _loc(f.get("so_what"), locale) if f.get("so_what") else ""
            caveats = "".join(f"<li>{esc(_loc(c, locale))}</li>" for c in f.get("caveats", []))
            details = "".join(f"<li>{esc(_loc(d, locale))}</li>" for d in f.get("details", [])[:6])
            items.append(
                f'<article class="finding"><div class="finding-head"><h3>{esc(_loc(f["title"], locale))}</h3>'
                f'<span class="conf {f["confidence"]}">{CONFIDENCE[f["confidence"]][locale]}</span></div>'
                f"{paragraphs(_loc(f['summary'], locale))}"
                + (
                    f'<p class="so-what"><strong>{_t("so_what", locale)}:</strong> {esc(so_what)}</p>'
                    if so_what
                    else ""
                )
                + (f'<div class="metrics">{metrics}</div>' if metrics else "")
                + (f'<ul class="details">{details}</ul>' if details else "")
                + (
                    f'<div class="chart">{svg_chart(f.get("chart"), locale)}</div>'
                    if f.get("chart")
                    else ""
                )
                + (
                    f"<details open><summary>{_t('caveats', locale)}</summary><ul>{caveats}</ul></details>"
                    if caveats
                    else ""
                )
                + (
                    f'<p class="evidence">{_t("evidence", locale)}: {esc(str(f.get("evidence_id") or "—"))}</p>'
                )
                + "</article>"
            )
        sections.append(
            f"<section><h2>{esc(_loc(section['title'], locale))}</h2>{''.join(items)}</section>"
        )
    questions = "".join(
        f"<li>{esc(_loc(q, locale))}</li>" for q in dossier.get("next_questions", [])
    )
    risks = "".join(f"<li>{esc(_loc(r, locale))}</li>" for r in dossier.get("risks", []))
    unverified = verification.get("unverified") or []
    method = (
        f"<p>{_t('engine', locale)}: {esc(str(engine.get('provider')))}"
        + (f" · {esc(str(engine.get('model')))}" if engine.get("model") else "")
        + "</p>"
        + f"<p>{_t('deterministic', locale)}</p>"
        + (
            f"<p>{_t('verification', locale)}: {verification.get('numbers_checked', 0)}"
            + (
                f" — {_t('unverified', locale)}: {esc(', '.join(unverified))}"
                if unverified
                else " ✓"
            )
            + "</p>"
            if verification.get("mode") != "deterministic"
            else ""
        )
    )
    generated = datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")
    font = (
        '"Noto Sans Arabic", "Noto Naskh Arabic", "DejaVu Sans", sans-serif'
        if rtl
        else '"Inter", "Helvetica Neue", "DejaVu Sans", sans-serif'
    )
    style = f"""
    body {{ font-family: {font}; color: #142536; margin: 0 auto; max-width: 960px; padding: 32px; line-height: 1.6; background: #fff; }}
    header {{ border-bottom: 3px solid #087f8c; margin-bottom: 24px; padding-bottom: 12px; }}
    header h1 {{ margin: 0; font-size: 28px; }} header p {{ margin: 4px 0; color: #52697e; }}
    .headline {{ font-size: 18px; font-weight: 600; color: #075e68; background: #eef8f7; padding: 12px 16px; border-radius: 10px; }}
    h2 {{ font-size: 20px; border-bottom: 1px solid #dfe4e3; padding-bottom: 6px; margin-top: 32px; }}
    h3 {{ font-size: 15px; margin: 0; }}
    .finding, .rec {{ border: 1px solid #dfe4e3; border-radius: 12px; padding: 14px 16px; margin: 12px 0; break-inside: avoid; }}
    .finding-head, .rec-head {{ display: flex; justify-content: space-between; gap: 12px; align-items: center; }}
    .conf {{ font-size: 11px; padding: 2px 8px; border-radius: 99px; white-space: nowrap; }}
    .conf.high {{ background: #dff3e8; color: #19704a; }} .conf.medium {{ background: #fcebc7; color: #99620a; }} .conf.low {{ background: #fde5e6; color: #a4363d; }}
    .badge {{ font-size: 11px; font-weight: 700; padding: 2px 8px; border-radius: 6px; color: #fff; background: #52697e; margin-inline-end: 8px; }}
    .badge.p1 {{ background: #a4363d; }} .badge.p2 {{ background: #99620a; }} .badge.p3 {{ background: #365e93; }}
    .rec-head {{ justify-content: flex-start; }}
    .metrics {{ display: flex; flex-wrap: wrap; gap: 8px; margin: 8px 0; }}
    .metric {{ background: #f8f9f7; border-radius: 8px; padding: 6px 10px; min-width: 120px; }}
    .metric span {{ display: block; font-size: 11px; color: #52697e; }} .metric strong {{ font-size: 15px; }}
    .so-what {{ background: #eef8f7; padding: 8px 10px; border-radius: 8px; }}
    .impact {{ color: #075e68; }} .meta, .evidence {{ font-size: 11px; color: #718397; }}
    details summary {{ font-size: 12px; color: #99620a; cursor: pointer; }} details li {{ font-size: 12px; color: #52697e; }}
    .chart {{ margin: 8px 0; }}
    footer {{ margin-top: 40px; font-size: 11px; color: #718397; border-top: 1px solid #dfe4e3; padding-top: 8px; }}
    """
    if for_pdf:
        style += "body { padding: 0; max-width: none; font-size: 10pt; } .finding, .rec { page-break-inside: avoid; }"
    title = f"{_t('title', locale)} — {dataset.get('name', '')}"
    return (
        f'<!doctype html><html lang="{locale}" dir="{"rtl" if rtl else "ltr"}"><head><meta charset="utf-8">'
        f"<title>{esc(title)}</title><style>{style}</style></head><body>"
        f"<header><h1>{esc(title)}</h1><p>{_t('prepared_by', locale)} · {generated}</p></header>"
        + (f'<p class="headline">{esc(headline)}</p>' if headline else "")
        + f"<section><h2>{_t('summary', locale)}</h2>{paragraphs(summary)}</section>"
        + f"<section><h2>{_t('recommendations', locale)}</h2>{''.join(recs)}</section>"
        + "".join(sections)
        + (f"<section><h2>{_t('risks', locale)}</h2><ul>{risks}</ul></section>" if risks else "")
        + (
            f"<section><h2>{_t('next', locale)}</h2><ul>{questions}</ul></section>"
            if questions
            else ""
        )
        + f"<section><h2>{_t('method', locale)}</h2>{method}</section>"
        + f"<footer>BASEERA · {esc(str(dataset.get('version_id', '')))}</footer></body></html>"
    )


def dossier_pdf(dossier: dict[str, Any], locale: str) -> bytes:
    from weasyprint import CSS, HTML
    from weasyprint.urls import URLFetcher

    fetcher = URLFetcher(allowed_protocols=(), allow_redirects=False, fail_on_errors=True)
    document = HTML(string=dossier_html(dossier, locale, for_pdf=True), url_fetcher=fetcher)
    stylesheet = CSS(
        string="@page { size: A4; margin: 16mm; @bottom-center { content: counter(page) ' / ' counter(pages); font-size: 9pt; } }",
        url_fetcher=fetcher,
    )
    return document.write_pdf(stylesheets=[stylesheet], pdf_tags=True)


# ---------------------------------------------------------------------------- Markdown
def dossier_markdown(dossier: dict[str, Any], locale: str) -> str:
    lines = [f"# {_t('title', locale)} — {dossier.get('dataset', {}).get('name', '')}", ""]
    if dossier.get("headline"):
        lines += [f"> {_loc(dossier['headline'], locale)}", ""]
    lines += [f"## {_t('summary', locale)}", "", _loc(dossier.get("executive_summary"), locale), ""]
    lines += [f"## {_t('recommendations', locale)}", ""]
    for rec in dossier.get("recommendations", []):
        lines.append(f"### [{rec['priority']}] {_loc(rec['title'], locale)}")
        lines.append(_loc(rec["rationale"], locale))
        lines.append(f"**{_t('impact', locale)}:** {_loc(rec['expected_impact'], locale)}")
        actions = (rec.get("actions_model") or {}).get(locale) or [
            _loc(a, locale) for a in rec.get("actions", [])
        ]
        lines += [f"- {a}" for a in actions] + [""]
    findings = {f["id"]: f for f in dossier.get("findings", [])}
    for section in dossier.get("sections", []):
        lines += [f"## {_loc(section['title'], locale)}", ""]
        for finding_id in section["findings"]:
            f = findings.get(finding_id)
            if not f:
                continue
            lines.append(f"### {_loc(f['title'], locale)} ({CONFIDENCE[f['confidence']][locale]})")
            lines.append(_loc(f["summary"], locale))
            for caveat in f.get("caveats", []):
                lines.append(f"- _{_loc(caveat, locale)}_")
            lines.append("")
    if dossier.get("next_questions"):
        lines += [f"## {_t('next', locale)}", ""] + [
            f"- {_loc(q, locale)}" for q in dossier["next_questions"]
        ]
    return "\n".join(lines) + "\n"


# -------------------------------------------------------------------------------- DOCX
def dossier_docx(dossier: dict[str, Any], locale: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    rtl = locale == "ar"
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Noto Sans Arabic" if rtl else "Calibri"
    style.font.size = Pt(10.5)

    def para(
        text: str, bold: bool = False, color: RGBColor | None = None, size: float | None = None
    ) -> None:
        p = document.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if color:
            run.font.color.rgb = color
        if size:
            run.font.size = Pt(size)
        if rtl:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.element.get_or_add_pPr().set(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi", "1"
            )

    def heading(text: str, level: int) -> None:
        h = document.add_heading(text, level=level)
        if rtl:
            h.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    heading(f"{_t('title', locale)} — {dossier.get('dataset', {}).get('name', '')}", 0)
    para(_t("prepared_by", locale), color=RGBColor(0x52, 0x69, 0x7E))
    if dossier.get("headline"):
        para(
            _loc(dossier["headline"], locale), bold=True, color=RGBColor(0x07, 0x5E, 0x68), size=13
        )
    heading(_t("summary", locale), 1)
    for block in _loc(dossier.get("executive_summary"), locale).split("\n\n"):
        para(block)
    heading(_t("recommendations", locale), 1)
    for rec in dossier.get("recommendations", []):
        heading(f"[{rec['priority']}] {_loc(rec['title'], locale)}", 2)
        para(_loc(rec["rationale"], locale))
        para(
            f"{_t('impact', locale)}: {_loc(rec['expected_impact'], locale)}",
            color=RGBColor(0x07, 0x5E, 0x68),
        )
        actions = (rec.get("actions_model") or {}).get(locale) or [
            _loc(a, locale) for a in rec.get("actions", [])
        ]
        for action in actions:
            p = document.add_paragraph(action, style="List Bullet")
            if rtl:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    findings = {f["id"]: f for f in dossier.get("findings", [])}
    for section in dossier.get("sections", []):
        heading(_loc(section["title"], locale), 1)
        for finding_id in section["findings"]:
            f = findings.get(finding_id)
            if not f:
                continue
            heading(f"{_loc(f['title'], locale)} — {CONFIDENCE[f['confidence']][locale]}", 2)
            for block in _loc(f["summary"], locale).split("\n\n"):
                para(block)
            metrics = f.get("metrics", [])[:4]
            if metrics:
                table = document.add_table(rows=2, cols=len(metrics))
                table.style = "Light Grid Accent 1"
                for i, m in enumerate(metrics):
                    table.cell(0, i).text = _loc(m.get("label"), locale)
                    table.cell(1, i).text = _metric_value(m)
            for caveat in f.get("caveats", []):
                para("• " + _loc(caveat, locale), color=RGBColor(0x99, 0x62, 0x0A), size=9)
    if dossier.get("next_questions"):
        heading(_t("next", locale), 1)
        for q in dossier["next_questions"]:
            document.add_paragraph(_loc(q, locale), style="List Bullet")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


# -------------------------------------------------------------------------------- PPTX
def dossier_pptx(dossier: dict[str, Any], locale: str) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    rtl = locale == "ar"
    align = PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT
    ink, teal, muted = (
        RGBColor(0x14, 0x25, 0x36),
        RGBColor(0x08, 0x7F, 0x8C),
        RGBColor(0x52, 0x69, 0x7E),
    )
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    blank = deck.slide_layouts[6]

    def text_box(
        slide: Any,
        text: str,
        left: float,
        top: float,
        width: float,
        height: float,
        size: int,
        bold: bool = False,
        color: Any = ink,
    ) -> None:
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        for index, block in enumerate(text.split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = align
            run = paragraph.add_run()
            run.text = block
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color

    def slide_with_title(title: str, eyebrow: str = "") -> Any:
        slide = deck.slides.add_slide(blank)
        bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = teal
        bar.line.fill.background()
        if eyebrow:
            text_box(slide, eyebrow, 0.6, 0.35, 12, 0.4, 12, color=teal)
        text_box(slide, title, 0.6, 0.7, 12, 0.9, 26, bold=True)
        return slide

    dataset = dossier.get("dataset", {})
    cover = slide_with_title(f"{_t('title', locale)}\n{dataset.get('name', '')}")
    text_box(cover, _t("prepared_by", locale), 0.6, 2.4, 12, 0.5, 14, color=muted)
    if dossier.get("headline"):
        text_box(
            cover, _loc(dossier["headline"], locale), 0.6, 3.2, 12, 1.5, 20, bold=True, color=teal
        )
    summary = slide_with_title(_t("summary", locale))
    text_box(
        summary, _loc(dossier.get("executive_summary"), locale)[:1800], 0.6, 1.7, 12.1, 5.5, 12
    )
    recs = dossier.get("recommendations", [])[:6]
    if recs:
        slide = slide_with_title(_t("recommendations", locale))
        body = "\n".join(
            f"[{r['priority']}] {_loc(r['title'], locale)} — {_loc(r['expected_impact'], locale)}"
            for r in recs
        )
        text_box(slide, body, 0.6, 1.7, 12.1, 5.5, 14)
    findings = [f for f in dossier.get("findings", []) if f["kind"] != "schema"][:10]
    for f in findings:
        slide = slide_with_title(_loc(f["title"], locale), CONFIDENCE[f["confidence"]][locale])
        chart = f.get("chart") or {}
        has_chart = chart.get("type") in {"line", "bar", "bar_horizontal", "pareto", "forecast"}
        text_box(
            slide, _loc(f["summary"], locale)[:900], 0.6, 1.7, 5.4 if has_chart else 12.1, 5.2, 13
        )
        if has_chart:
            data = CategoryChartData()
            if chart["type"] == "forecast":
                categories = [str(x) for x in chart.get("x", [])]
                history = chart.get("history", [])
                future = chart.get("forecast", [])
                data.categories = categories
                data.add_series("history", [*history, *([None] * len(future))])
                data.add_series("forecast", [*([None] * len(history)), *future])
                kind = XL_CHART_TYPE.LINE
            elif chart["type"] == "line":
                data.categories = [str(x) for x in chart.get("x", [])]
                for series in chart.get("series", [])[:4]:
                    data.add_series(
                        str(series.get("name")),
                        [v if isinstance(v, int | float) else None for v in series.get("data", [])],
                    )
                kind = XL_CHART_TYPE.LINE
            else:
                raw_labels = chart.get("x_ar") if rtl and chart.get("x_ar") else chart.get("x")
                data.categories = [str(x) for x in raw_labels or []][:20]
                first = (chart.get("series") or [{}])[0]
                data.add_series(
                    str(first.get("name", "")),
                    [v if isinstance(v, int | float) else 0 for v in first.get("data", [])][:20],
                )
                kind = (
                    XL_CHART_TYPE.BAR_CLUSTERED
                    if chart["type"] == "bar_horizontal"
                    else XL_CHART_TYPE.COLUMN_CLUSTERED
                )
            try:
                graphic = slide.shapes.add_chart(
                    kind, Inches(6.3), Inches(1.7), Inches(6.5), Inches(5.2), data
                )
                graphic.chart.has_legend = chart["type"] in {"forecast", "line"}
                if graphic.chart.has_legend:
                    graphic.chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                    graphic.chart.legend.include_in_layout = False
            except (ValueError, TypeError):
                pass
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def export(dossier: dict[str, Any], kind: str, locale: str) -> tuple[bytes, str, str]:
    if kind not in MEDIA_TYPES:
        raise AppError(422, "unsupported_export", f"Supported formats: {', '.join(MEDIA_TYPES)}.")
    media_type, extension = MEDIA_TYPES[kind]
    if kind == "html":
        content = dossier_html(dossier, locale).encode("utf-8")
    elif kind == "pdf":
        content = dossier_pdf(dossier, locale)
    elif kind == "docx":
        content = dossier_docx(dossier, locale)
    elif kind == "pptx":
        content = dossier_pptx(dossier, locale)
    elif kind == "md":
        content = dossier_markdown(dossier, locale).encode("utf-8")
    else:
        content = json.dumps(dossier, ensure_ascii=False, indent=2, default=str).encode("utf-8")
    return content, media_type, extension
