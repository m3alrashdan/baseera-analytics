from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
import tempfile
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import pyarrow.parquet as pq

from .cleaning import control_totals, run_recipe
from .errors import AppError
from .profiling import as_number, profile_dataset

SUPPORTED_FORMATS = {"csv", "tsv", "json", "jsonl", "ndjson", "xlsx", "parquet"}
IDENTIFIER_NAMES = {"id", "identifier", "order_id", "customer_id", "source_id"}
AMBIGUOUS_DATE = re.compile(r"^\d{1,2}/\d{1,2}/\d{4}$")


def canonical_json(value: Any) -> str:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=_json
    )


def stable_hash(value: Any) -> str:
    return hashlib.sha256(canonical_json(value).encode("utf-8")).hexdigest()


def _json(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if hasattr(value, "item"):
        return value.item()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    raise TypeError(f"Unsupported value: {type(value)!r}")


class ArtifactStore:
    """Tenant-namespaced local object store with atomic writes and path containment."""

    def __init__(self, root: Path) -> None:
        self.root = root.resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def _resolve(self, object_key: str) -> Path:
        if not object_key or Path(object_key).is_absolute():
            raise ValueError("Object keys must be relative")
        target = (self.root / object_key).resolve()
        if self.root not in target.parents:
            raise ValueError("Object key escapes artifact root")
        return target

    def write_bytes(self, object_key: str, content: bytes) -> None:
        target = self._resolve(object_key)
        target.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary = tempfile.mkstemp(prefix=".baseera-", dir=target.parent)
        try:
            with os.fdopen(descriptor, "wb") as handle:
                handle.write(content)
                handle.flush()
                os.fsync(handle.fileno())
            Path(temporary).replace(target)
        finally:
            Path(temporary).unlink(missing_ok=True)

    def read_bytes(self, object_key: str) -> bytes:
        target = self._resolve(object_key)
        try:
            return target.read_bytes()
        except FileNotFoundError as exc:
            # A record can outlive its bytes: a restored database pointed at a different
            # artifact root, or a partially copied backup. Say so plainly instead of
            # surfacing a stack trace from deep inside a request handler.
            raise AppError(
                410,
                "artifact_missing",
                "The stored data for this version is not present in the artifact store. "
                "The record still exists, but its rows cannot be read. Check that the "
                "artifact root matches the database being used.",
                details={"object_key": object_key},
            ) from exc

    def write_rows(self, object_key: str, rows: list[dict[str, Any]]) -> None:
        self.write_bytes(object_key, canonical_json(rows).encode("utf-8"))

    def read_rows(self, object_key: str) -> list[dict[str, Any]]:
        payload = json.loads(self.read_bytes(object_key))
        if not isinstance(payload, list):
            raise ValueError("Dataset artifact is not a row collection")
        return payload


@dataclass(slots=True)
class ParsedUpload:
    rows: list[dict[str, Any]]
    columns: list[str]
    rows_discovered: int
    rejected: list[dict[str, Any]] = field(default_factory=list)
    extraction: dict[str, Any] = field(default_factory=dict)

    @property
    def coverage(self) -> dict[str, int]:
        accepted = len(self.rows)
        return {
            "rows_discovered": self.rows_discovered,
            "rows_parsed": accepted,
            "rows_accepted": accepted,
            "rows_rejected": len(self.rejected),
            "rows_sampled": accepted,
            "rows_analyzed": accepted,
        }


def infer_format(filename: str) -> str:
    suffix = Path(filename).suffix.lower().lstrip(".")
    if suffix == "xls":
        raise AppError(
            415,
            "unsupported_format",
            "Legacy XLS is unsupported; convert the file to XLSX or CSV.",
        )
    if suffix not in SUPPORTED_FORMATS:
        raise AppError(415, "unsupported_format", f"Unsupported file format: {suffix or 'unknown'}")
    return "jsonl" if suffix == "ndjson" else suffix


def parse_upload(filename: str, content: bytes) -> tuple[str, ParsedUpload]:
    kind = infer_format(filename)
    try:
        if kind in {"csv", "tsv"}:
            parsed = _parse_delimited(content, "\t" if kind == "tsv" else ",")
        elif kind == "json":
            parsed = _parse_json(content)
        elif kind == "jsonl":
            parsed = _parse_jsonl(content)
        elif kind == "xlsx":
            parsed = _parse_xlsx(content)
        else:
            parsed = _parse_parquet(content)
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            422,
            "parse_failed",
            "The file could not be parsed safely.",
            details={"format": kind, "reason": type(exc).__name__},
        ) from exc
    parsed.extraction = {
        "format": kind,
        "scope": "full",
        "content_bytes": len(content),
        **parsed.extraction,
    }
    return kind, parsed


def _validate_columns(columns: list[str]) -> list[str]:
    cleaned = [str(value).strip() for value in columns]
    if not cleaned or any(not value for value in cleaned):
        raise AppError(422, "invalid_schema", "Every column must have a non-empty name")
    if len(set(cleaned)) != len(cleaned):
        raise AppError(422, "invalid_schema", "Duplicate column names are not accepted")
    return cleaned


def _parse_delimited(content: bytes, delimiter: str) -> ParsedUpload:
    text = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw = list(reader)
    if not raw:
        raise AppError(422, "empty_file", "The file contains no header")
    columns = _validate_columns(raw[0])
    rows: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []
    for line_number, values in enumerate(raw[1:], start=2):
        if len(values) != len(columns):
            rejected.append(
                {
                    "source_location": f"line:{line_number}",
                    "error_code": "column_count_mismatch",
                    "raw_payload": {"values": values},
                }
            )
            continue
        rows.append(
            {
                column: (value if value != "" else None)
                for column, value in zip(columns, values, strict=True)
            }
        )
    return ParsedUpload(rows, columns, len(raw) - 1, rejected)


def _parse_json(content: bytes) -> ParsedUpload:
    value = json.loads(content.decode("utf-8-sig"))
    if not isinstance(value, list):
        raise AppError(
            422, "invalid_json_shape", "JSON uploads must contain a top-level record array"
        )
    return _records_to_parsed(value, "index")


def _parse_jsonl(content: bytes) -> ParsedUpload:
    records: list[Any] = []
    rejected: list[dict[str, Any]] = []
    lines = [line for line in content.decode("utf-8-sig").splitlines() if line.strip()]
    for line_number, line in enumerate(lines, start=1):
        try:
            records.append(json.loads(line))
        except json.JSONDecodeError:
            rejected.append(
                {
                    "source_location": f"line:{line_number}",
                    "error_code": "invalid_json",
                    "raw_payload": {"text": line[:500]},
                }
            )
    parsed = _records_to_parsed(records, "line")
    parsed.rows_discovered = len(lines)
    parsed.rejected.extend(rejected)
    return parsed


def _records_to_parsed(records: list[Any], location: str) -> ParsedUpload:
    rows: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []
    columns: list[str] = []
    for index, value in enumerate(records, start=1):
        if not isinstance(value, dict) or not all(isinstance(key, str) for key in value):
            rejected.append(
                {
                    "source_location": f"{location}:{index}",
                    "error_code": "record_not_object",
                    "raw_payload": {"value": value},
                }
            )
            continue
        for key in value:
            if key not in columns:
                columns.append(key)
        rows.append({key: _json_safe(item) for key, item in value.items()})
    columns = _validate_columns(columns) if columns else []
    normalized = [{column: row.get(column) for column in columns} for row in rows]
    return ParsedUpload(normalized, columns, len(records), rejected)


def _parse_xlsx(content: bytes) -> ParsedUpload:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(entry.file_size for entry in entries) > 128 * 1024 * 1024:
            raise AppError(
                413,
                "extraction_limit_exceeded",
                "Workbook expanded size exceeds the safe extraction budget",
            )
    workbook = openpyxl.load_workbook(
        io.BytesIO(content), read_only=True, data_only=True, keep_links=False
    )
    formulas = openpyxl.load_workbook(
        io.BytesIO(content), read_only=True, data_only=False, keep_links=False
    )
    try:
        inventory = [
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
            }
            for sheet in workbook.worksheets
        ]
        selected = None
        values: list[Any] = []
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue
            if (
                (worksheet.max_row or 0) > 100_001
                or (worksheet.max_column or 0) > 500
                or (worksheet.max_row or 0) * (worksheet.max_column or 0) > 2_000_000
            ):
                raise AppError(
                    413,
                    "extraction_limit_exceeded",
                    "Workbook sheet exceeds the row or cell extraction budget",
                )
            values = list(worksheet.iter_rows(values_only=True))
            if any(any(value is not None for value in row) for row in values):
                selected = worksheet.title
                break
        if selected is None:
            raise AppError(422, "empty_workbook", "The workbook has no non-empty visible sheets")
        while values and all(value is None for value in values[0]):
            values.pop(0)
        columns = _validate_columns([str(value or "") for value in values[0]])
        records = [
            {column: _json_safe(value) for column, value in zip(columns, row, strict=True)}
            for row in values[1:]
        ]
        formula_count = 0
        missing_cache = 0
        cached_rows = workbook[selected].iter_rows(values_only=True)
        for cached, formula_row in zip(cached_rows, formulas[selected].iter_rows(), strict=True):
            for value, cell in zip(cached, formula_row, strict=True):
                if cell.data_type == "f":
                    formula_count += 1
                    missing_cache += value is None
        return ParsedUpload(
            records,
            columns,
            len(records),
            extraction={
                "sheets": workbook.sheetnames,
                "selected_sheet": selected,
                "sheet_selection_policy": "first_visible_non_empty",
                "sheet_inventory": inventory,
                "formulas_evaluated": False,
                "formula_evaluation": "cached_values_only",
                "formula_cells_detected": formula_count,
                "formula_cells_without_cached_value": missing_cache,
                "limits": {
                    "rows": 100000,
                    "columns": 500,
                    "cells": 2000000,
                    "expanded_bytes": 134217728,
                },
            },
        )
    finally:
        workbook.close()
        formulas.close()


def _parse_parquet(content: bytes) -> ParsedUpload:
    table = pq.read_table(io.BytesIO(content))
    columns = _validate_columns(table.column_names)
    records = [{key: _json_safe(value) for key, value in row.items()} for row in table.to_pylist()]
    return ParsedUpload(
        records,
        columns,
        len(records),
        extraction={
            "row_groups": table.num_rows and pq.ParquetFile(io.BytesIO(content)).num_row_groups
        },
    )


def _json_safe(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if hasattr(value, "item"):
        return _json_safe(value.item())
    return str(value)


def profile_rows(rows: list[dict[str, Any]], columns: list[str]) -> dict[str, Any]:
    """Profile a dataset version.

    Delegates to :mod:`baseera.profiling`, which returns the historical keys plus the
    per-column statistics, format anomalies, quality score and suggested repairs the
    review workspace needs.
    """
    return profile_dataset(rows, columns)


def apply_cleaning_recipe(
    source_rows: list[dict[str, Any]], columns: list[str], recipe: dict[str, Any]
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Run a reviewed recipe and return the cleaned rows with the review preview.

    Kept for callers that do not need the post-recipe column list; steps that add or
    remove columns are better served by :func:`baseera.cleaning.run_recipe` directly.
    """
    rows, _columns, preview = run_recipe(source_rows, columns, recipe)
    return rows, preview


def _as_number(value: Any) -> float | None:
    return as_number(value)


def _control_total(rows: list[dict[str, Any]], column: str) -> float:
    return (control_totals(rows, [column]).get(column) or {}).get("sum", 0.0)
